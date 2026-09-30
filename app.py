from __future__ import annotations

import csv
import hashlib
import io
import os
from datetime import datetime
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, Request, Form, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy import (
    create_engine, String, Integer, DateTime, ForeignKey, LargeBinary,
    select, func, case, cast, delete, update, inspect, text
)
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker

BASE_DIR = Path(__file__).resolve().parent
LOCAL_DB = BASE_DIR / "data" / "eleccion.db"
LOCAL_DB.parent.mkdir(parents=True, exist_ok=True)

raw_db_url = os.environ.get("DATABASE_URL", "").strip()
if raw_db_url:
    if raw_db_url.startswith("postgres://"):
        raw_db_url = "postgresql+psycopg://" + raw_db_url[len("postgres://"):]
    elif raw_db_url.startswith("postgresql://"):
        raw_db_url = "postgresql+psycopg://" + raw_db_url[len("postgresql://"):]
    DATABASE_URL = raw_db_url
else:
    DATABASE_URL = f"sqlite:///{LOCAL_DB.as_posix()}"

engine_kwargs = {"pool_pre_ping": True}
if DATABASE_URL.startswith("sqlite"):
    engine_kwargs["connect_args"] = {"check_same_thread": False}
engine = create_engine(DATABASE_URL, **engine_kwargs)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)


class Base(DeclarativeBase):
    pass


class Setting(Base):
    __tablename__ = "settings"
    key: Mapped[str] = mapped_column(String(100), primary_key=True)
    value: Mapped[str] = mapped_column(String(2000), nullable=False)


class Student(Base):
    __tablename__ = "students"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    # El DNI es el identificador de acceso del elector. Grado y sección se mantienen
    # como campos opcionales para reportes administrativos, pero nunca se muestran
    # ni se usan para identificar al estudiante en la pantalla pública.
    dni: Mapped[Optional[str]] = mapped_column(String(20), nullable=True)
    grade: Mapped[str] = mapped_column(String(20), nullable=False, default="", index=True)
    section: Mapped[str] = mapped_column(String(20), nullable=False, default="", index=True)
    full_name: Mapped[str] = mapped_column(String(300), nullable=False)
    voted_at: Mapped[Optional[datetime]] = mapped_column(DateTime, nullable=True)


class Candidate(Base):
    __tablename__ = "candidates"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    list_number: Mapped[str] = mapped_column(String(30), nullable=False)
    list_name: Mapped[str] = mapped_column(String(200), nullable=False)
    candidate_name: Mapped[str] = mapped_column(String(300), nullable=False)
    symbol_data: Mapped[Optional[bytes]] = mapped_column(LargeBinary, nullable=True)
    symbol_mime: Mapped[Optional[str]] = mapped_column(String(80), nullable=True)
    active: Mapped[int] = mapped_column(Integer, nullable=False, default=1)
    created_at: Mapped[datetime] = mapped_column(DateTime, nullable=False, default=datetime.now)

    @property
    def symbol_path(self):
        return f"candidate-symbol/{self.id}" if self.symbol_data else None


class Vote(Base):
    __tablename__ = "votes"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    candidate_id: Mapped[Optional[int]] = mapped_column(ForeignKey("candidates.id"), nullable=True, index=True)
    cast_at: Mapped[datetime] = mapped_column(DateTime, nullable=False, default=datetime.now)


app = FastAPI(title="Elecciones Escolares FBC")
app.add_middleware(
    SessionMiddleware,
    secret_key=os.environ.get("ELECCIONES_SECRET", "fbc-local-secret-cambiar-en-produccion"),
    same_site="lax",
    https_only=os.environ.get("RENDER", "").lower() == "true",
)
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


def now_dt():
    return datetime.now()


def natural_number_key(value):
    """Ordena valores numericos primero y deja texto como respaldo, sin CAST en PostgreSQL."""
    text = str(value or "").strip()
    try:
        return (0, int(text), text.lower())
    except ValueError:
        return (1, 0, text.lower())



def normalize_dni(value: str) -> str:
    return "".join(ch for ch in str(value or "").strip() if ch.isdigit())


def ensure_student_dni_schema():
    """Migra de forma segura la versión anterior sin borrar padrón/candidatos/votos."""
    inspector = inspect(engine)
    if "students" not in inspector.get_table_names():
        return
    columns = {c["name"] for c in inspector.get_columns("students")}
    with engine.begin() as conn:
        if "dni" not in columns:
            conn.execute(text("ALTER TABLE students ADD COLUMN dni VARCHAR(20)"))
        # PostgreSQL y SQLite aceptan índices únicos con múltiples NULL.
        conn.execute(text("CREATE UNIQUE INDEX IF NOT EXISTS ux_students_dni ON students (dni)"))


def password_hash(password: str) -> str:
    return hashlib.sha256(password.encode("utf-8")).hexdigest()


def init_db():
    Base.metadata.create_all(engine)
    ensure_student_dni_schema()
    defaults = {
        "school_name": "I.E. Francisco Bolognesi Cervantes",
        "election_title": "Elecciones Municipales Escolares 2026",
        "election_open": "0",
        "show_results_live": "0",
        "admin_password_hash": password_hash("FBC2026!"),
        "first_run": "1",
    }
    with SessionLocal.begin() as db:
        existing = set(db.scalars(select(Setting.key)).all())
        for k, v in defaults.items():
            if k not in existing:
                db.add(Setting(key=k, value=v))


def setting(key: str, default: str = "") -> str:
    with SessionLocal() as db:
        row = db.get(Setting, key)
        return row.value if row else default


def set_setting(key: str, value: str):
    with SessionLocal.begin() as db:
        row = db.get(Setting, key)
        if row:
            row.value = value
        else:
            db.add(Setting(key=key, value=value))


def base_context(request: Request, **extra):
    ctx = {
        "request": request,
        "school_name": setting("school_name"),
        "election_title": setting("election_title"),
        "election_open": setting("election_open") == "1",
        "is_admin": bool(request.session.get("admin")),
    }
    ctx.update(extra)
    return ctx


def admin_required(request: Request):
    return bool(request.session.get("admin"))


init_db()


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    with SessionLocal() as db:
        total = db.scalar(select(func.count(Student.id))) or 0
        candidates = db.scalar(select(func.count(Candidate.id)).where(Candidate.active == 1)) or 0
    return templates.TemplateResponse(
        request=request,
        name="home.html",
        context=base_context(request, total_students=total, candidate_count=candidates),
    )


@app.post("/identify")
def identify(request: Request, dni: str = Form(...)):
    if setting("election_open") != "1":
        return RedirectResponse("/?msg=La+votación+está+cerrada", status_code=303)

    dni_clean = normalize_dni(dni)
    if len(dni_clean) != 8:
        return RedirectResponse("/?msg=Ingresa+un+DNI+válido+de+8+dígitos", status_code=303)

    with SessionLocal() as db:
        row = db.scalar(select(Student).where(Student.dni == dni_clean))
        if not row:
            return RedirectResponse("/?msg=No+se+pudo+validar+el+DNI+ingresado", status_code=303)
        if row.voted_at:
            return RedirectResponse("/?msg=Este+DNI+ya+registró+su+participación", status_code=303)
        request.session["student_id"] = row.id
        request.session.pop("selected_candidate", None)
    return RedirectResponse("/vote", status_code=303)


@app.get("/vote", response_class=HTMLResponse)
def vote_page(request: Request):
    sid = request.session.get("student_id")
    if not sid:
        return RedirectResponse("/", status_code=303)
    if setting("election_open") != "1":
        request.session.pop("student_id", None)
        return RedirectResponse("/?msg=La+votación+está+cerrada", status_code=303)
    with SessionLocal() as db:
        student = db.get(Student, sid)
        candidates = db.scalars(
            select(Candidate).where(Candidate.active == 1).order_by(Candidate.list_number, Candidate.id)
        ).all()
        candidates.sort(key=lambda c: (natural_number_key(c.list_number), c.id))
        if not student or student.voted_at:
            request.session.pop("student_id", None)
            return RedirectResponse("/?msg=Este+estudiante+ya+votó", status_code=303)
        db.expunge(student)
        for c in candidates:
            db.expunge(c)
    return templates.TemplateResponse(request=request, name="vote.html", context=base_context(request, student=student, candidates=candidates))


@app.post("/vote/select", response_class=HTMLResponse)
def select_vote(request: Request, candidate_id: str = Form(...)):
    sid = request.session.get("student_id")
    if not sid:
        return RedirectResponse("/", status_code=303)
    is_blank = candidate_id == "blank"
    with SessionLocal() as db:
        student = db.get(Student, sid)
        candidate = None if is_blank else db.get(Candidate, int(candidate_id))
        if not student or student.voted_at:
            request.session.pop("student_id", None)
            return RedirectResponse("/?msg=Este+estudiante+ya+votó", status_code=303)
        if not is_blank and (not candidate or not candidate.active):
            return RedirectResponse("/vote", status_code=303)
        request.session["selected_candidate"] = "blank" if is_blank else int(candidate.id)
        db.expunge(student)
        if candidate:
            db.expunge(candidate)
    return templates.TemplateResponse(
        request=request,
        name="confirm.html",
        context=base_context(request, student=student, candidate=candidate, is_blank=is_blank),
    )


@app.post("/vote/confirm", response_class=HTMLResponse)
def confirm_vote(request: Request):
    sid = request.session.get("student_id")
    selected = request.session.get("selected_candidate")
    if not sid or selected is None:
        return RedirectResponse("/", status_code=303)
    if setting("election_open") != "1":
        request.session.pop("student_id", None)
        request.session.pop("selected_candidate", None)
        return RedirectResponse("/?msg=La+votación+está+cerrada", status_code=303)

    try:
        with SessionLocal.begin() as db:
            # En PostgreSQL bloquea la fila del estudiante durante la transacción.
            student = db.scalar(select(Student).where(Student.id == sid).with_for_update())
            if not student or student.voted_at:
                request.session.pop("student_id", None)
                request.session.pop("selected_candidate", None)
                return RedirectResponse("/?msg=Este+estudiante+ya+registró+su+participación", status_code=303)

            candidate_id = None if selected == "blank" else int(selected)
            if candidate_id is not None:
                candidate = db.get(Candidate, candidate_id)
                if not candidate or not candidate.active:
                    return RedirectResponse("/vote", status_code=303)

            db.add(Vote(candidate_id=candidate_id, cast_at=now_dt()))
            student.voted_at = now_dt()
    except Exception:
        request.session.pop("student_id", None)
        request.session.pop("selected_candidate", None)
        raise

    request.session.pop("student_id", None)
    request.session.pop("selected_candidate", None)
    return templates.TemplateResponse(request=request, name="thanks.html", context=base_context(request))


@app.get("/candidate-symbol/{candidate_id}")
def candidate_symbol(candidate_id: int):
    with SessionLocal() as db:
        c = db.get(Candidate, candidate_id)
        if not c or not c.symbol_data:
            raise HTTPException(status_code=404)
        return Response(content=c.symbol_data, media_type=c.symbol_mime or "image/png", headers={"Cache-Control": "public, max-age=3600"})


@app.get("/admin/login", response_class=HTMLResponse)
def admin_login_page(request: Request):
    return templates.TemplateResponse(request=request, name="admin_login.html", context=base_context(request))


@app.post("/admin/login")
def admin_login(request: Request, password: str = Form(...)):
    if password_hash(password) == setting("admin_password_hash"):
        request.session["admin"] = True
        return RedirectResponse("/admin", status_code=303)
    return RedirectResponse("/admin/login?error=1", status_code=303)


@app.get("/admin/logout")
def admin_logout(request: Request):
    request.session.clear()
    return RedirectResponse("/", status_code=303)


def dashboard_data():
    with SessionLocal() as db:
        total = db.scalar(select(func.count(Student.id))) or 0
        voted = db.scalar(select(func.count(Student.id)).where(Student.voted_at.is_not(None))) or 0
        turnout = round((voted / total * 100), 1) if total else 0

        classrooms = db.execute(
            select(
                Student.grade,
                Student.section,
                func.count(Student.id).label("total"),
                func.sum(case((Student.voted_at.is_not(None), 1), else_=0)).label("voted"),
            )
            .where(Student.grade != "", Student.section != "")
            .group_by(Student.grade, Student.section)
            .order_by(Student.grade, Student.section)
        ).all()

        candidates = db.scalars(
            select(Candidate).where(Candidate.active == 1).order_by(Candidate.list_number, Candidate.id)
        ).all()
        candidates.sort(key=lambda c: (natural_number_key(c.list_number), c.id))
        vote_counts = dict(db.execute(
            select(Vote.candidate_id, func.count(Vote.id)).where(Vote.candidate_id.is_not(None)).group_by(Vote.candidate_id)
        ).all())
        blank = db.scalar(select(func.count(Vote.id)).where(Vote.candidate_id.is_(None))) or 0
        total_votes = db.scalar(select(func.count(Vote.id))) or 0

        results = []
        for c in candidates:
            votes = int(vote_counts.get(c.id, 0))
            results.append({
                "id": c.id,
                "list_number": c.list_number,
                "list_name": c.list_name,
                "candidate_name": c.candidate_name,
                "symbol_path": c.symbol_path,
                "votes": votes,
                "pct": round((votes / total_votes * 100), 1) if total_votes else 0,
            })
        results.sort(key=lambda x: (-x["votes"], natural_number_key(x["list_number"])))

        classrooms = sorted(classrooms, key=lambda r: (natural_number_key(r.grade), str(r.section).lower()))
        classrooms_out = []
        for r in classrooms:
            v = int(r.voted or 0)
            t = int(r.total or 0)
            classrooms_out.append({
                "grade": r.grade,
                "section": r.section,
                "total": t,
                "voted": v,
                "missing": t - v,
                "pct": round((v / t * 100), 1) if t else 0,
            })

    blank_result = {
        "list_number": "—",
        "list_name": "Voto en blanco",
        "candidate_name": "",
        "symbol_path": None,
        "votes": blank,
        "pct": round((blank / total_votes * 100), 1) if total_votes else 0,
    }
    return total, voted, turnout, classrooms_out, results, blank_result, total_votes


@app.get("/admin", response_class=HTMLResponse)
def admin_dashboard(request: Request):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    total, voted, turnout, classrooms, results, blank_result, total_votes = dashboard_data()
    with SessionLocal() as db:
        candidates = db.scalars(select(Candidate).order_by(Candidate.list_number, Candidate.id)).all()
        candidates.sort(key=lambda c: (natural_number_key(c.list_number), c.id))
        for c in candidates:
            db.expunge(c)
    show_results = setting("show_results_live") == "1" or setting("election_open") != "1"
    return templates.TemplateResponse(
        request=request,
        name="admin.html",
        context=base_context(
            request,
            total=total,
            voted=voted,
            missing=total - voted,
            turnout=turnout,
            classrooms=classrooms,
            results=results,
            blank_result=blank_result,
            total_votes=total_votes,
            candidates=candidates,
            show_results=show_results,
            show_results_live=setting("show_results_live") == "1",
            first_run=setting("first_run") == "1",
        ),
    )


@app.post("/admin/settings")
def admin_settings(
    request: Request,
    school_name: str = Form(...),
    election_title: str = Form(...),
    show_results_live: Optional[str] = Form(None),
):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    set_setting("school_name", school_name.strip())
    set_setting("election_title", election_title.strip())
    set_setting("show_results_live", "1" if show_results_live else "0")
    set_setting("first_run", "0")
    return RedirectResponse("/admin#configuracion", status_code=303)


@app.post("/admin/password")
def admin_password(request: Request, current_password: str = Form(...), new_password: str = Form(...)):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    if password_hash(current_password) != setting("admin_password_hash"):
        return RedirectResponse("/admin?msg=Contraseña+actual+incorrecta#configuracion", status_code=303)
    if len(new_password) < 8:
        return RedirectResponse("/admin?msg=La+nueva+contraseña+debe+tener+al+menos+8+caracteres#configuracion", status_code=303)
    set_setting("admin_password_hash", password_hash(new_password))
    return RedirectResponse("/admin?msg=Contraseña+actualizada#configuracion", status_code=303)


@app.post("/admin/toggle-election")
def admin_toggle_election(request: Request, action: str = Form(...)):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    set_setting("election_open", "1" if action == "open" else "0")
    return RedirectResponse("/admin", status_code=303)


@app.post("/admin/candidates/add")
async def admin_candidate_add(
    request: Request,
    list_number: str = Form(...),
    list_name: str = Form(...),
    candidate_name: str = Form(...),
    symbol: Optional[UploadFile] = File(None),
):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)

    symbol_data = None
    symbol_mime = None
    if symbol and symbol.filename:
        ext = Path(symbol.filename).suffix.lower()
        if ext not in {".png", ".jpg", ".jpeg", ".webp"}:
            return RedirectResponse("/admin?msg=Formato+de+imagen+no+permitido#candidatos", status_code=303)
        symbol_data = await symbol.read()
        if len(symbol_data) > 4 * 1024 * 1024:
            return RedirectResponse("/admin?msg=La+imagen+supera+4MB#candidatos", status_code=303)
        symbol_mime = symbol.content_type or {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}[ext]

    with SessionLocal.begin() as db:
        db.add(Candidate(
            list_number=list_number.strip(),
            list_name=list_name.strip(),
            candidate_name=candidate_name.strip(),
            symbol_data=symbol_data,
            symbol_mime=symbol_mime,
            created_at=now_dt(),
            active=1,
        ))
    return RedirectResponse("/admin#candidatos", status_code=303)


@app.post("/admin/candidates/{candidate_id}/toggle")
def admin_candidate_toggle(candidate_id: int, request: Request):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    with SessionLocal.begin() as db:
        row = db.get(Candidate, candidate_id)
        if row:
            row.active = 0 if row.active else 1
    return RedirectResponse("/admin#candidatos", status_code=303)


@app.post("/admin/roster/upload")
async def admin_roster_upload(request: Request, roster: UploadFile = File(...), mode: str = Form("append")):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    filename = roster.filename or ""
    data = await roster.read()
    records = []

    try:
        if filename.lower().endswith(".csv"):
            text_data = data.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text_data))
            for row in reader:
                dni = normalize_dni(row.get("dni") or row.get("DNI") or "")
                name = str(row.get("apellidos_nombres") or row.get("nombre") or row.get("nombres") or row.get("full_name") or "").strip()
                grade = str(row.get("grado") or row.get("grade") or "").strip()
                section = str(row.get("seccion") or row.get("sección") or row.get("section") or "").strip().upper()
                if dni and name:
                    if len(dni) != 8:
                        raise ValueError(f"DNI inválido: {dni}")
                    records.append((dni, name, grade, section))
        elif filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.active
            headers = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
            idx = {h: i for i, h in enumerate(headers)}

            def find_index(options):
                for o in options:
                    if o in idx:
                        return idx[o]
                return None

            di = find_index(["dni", "documento", "numero_dni", "número_dni"])
            ni = find_index(["apellidos_nombres", "nombre", "nombres", "full_name", "apellidos y nombres"])
            gi = find_index(["grado", "grade"])
            si = find_index(["seccion", "sección", "section"])
            if di is None or ni is None:
                raise ValueError("Faltan columnas dni y/o apellidos_nombres")

            for row in ws.iter_rows(min_row=2, values_only=True):
                raw_dni = row[di] if di < len(row) else ""
                dni = normalize_dni(raw_dni)
                name = str(row[ni] or "").strip() if ni < len(row) else ""
                grade = str(row[gi] or "").strip() if gi is not None and gi < len(row) else ""
                section = str(row[si] or "").strip().upper() if si is not None and si < len(row) else ""
                if dni and name:
                    if len(dni) != 8:
                        raise ValueError(f"DNI inválido: {dni}")
                    records.append((dni, name, grade, section))
        else:
            return RedirectResponse("/admin?msg=Solo+se+admite+CSV+o+XLSX#padron", status_code=303)
    except Exception as exc:
        return RedirectResponse("/admin?msg=No+se+pudo+leer+el+archivo.+Verifica+que+incluya+DNI+y+apellidos_nombres#padron", status_code=303)

    if not records:
        return RedirectResponse("/admin?msg=No+se+encontraron+estudiantes+válidos#padron", status_code=303)

    # DNI único: si aparece repetido en el archivo, conserva la primera aparición.
    unique = {}
    for dni, name, grade, section in records:
        unique.setdefault(dni, (dni, name, grade, section))
    unique_records = list(unique.values())

    with SessionLocal.begin() as db:
        if mode == "replace":
            votes = db.scalar(select(func.count(Vote.id))) or 0
            if votes:
                return RedirectResponse("/admin?msg=No+se+puede+reemplazar+el+padrón+si+ya+existen+votos#padron", status_code=303)
            db.execute(delete(Student))

        existing_dnis = set(d for d in db.scalars(select(Student.dni).where(Student.dni.is_not(None))).all() if d)
        added = 0
        for dni, name, grade, section in unique_records:
            if dni not in existing_dnis:
                db.add(Student(dni=dni, full_name=name, grade=grade, section=section))
                existing_dnis.add(dni)
                added += 1

    return RedirectResponse(f"/admin?msg=Se+cargaron+{added}+estudiantes#padron", status_code=303)


@app.post("/admin/reset")
def admin_reset(request: Request, confirmation: str = Form(...)):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    if confirmation.strip().upper() != "REINICIAR":
        return RedirectResponse("/admin?msg=Escribe+REINICIAR+para+confirmar#reinicio", status_code=303)
    with SessionLocal.begin() as db:
        db.execute(delete(Vote))
        db.execute(update(Student).values(voted_at=None))
    set_setting("election_open", "0")
    return RedirectResponse("/admin?msg=Votación+reiniciada.+El+padrón+y+los+candidatos+se+conservaron#reinicio", status_code=303)


@app.get("/admin/export/turnout.csv")
def export_turnout(request: Request):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    with SessionLocal() as db:
        rows = db.scalars(select(Student).order_by(Student.full_name)).all()
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["dni", "apellidos_nombres", "grado", "seccion", "participo", "fecha_hora_voto"])
    for r in rows:
        w.writerow([r.dni or "", r.full_name, r.grade or "", r.section or "", "SI" if r.voted_at else "NO", r.voted_at.isoformat(sep=" ", timespec="seconds") if r.voted_at else ""])
    data = out.getvalue().encode("utf-8-sig")
    return StreamingResponse(io.BytesIO(data), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=padron_participacion_elecciones.csv"})


@app.get("/admin/export/turnout.xlsx")
def export_turnout_xlsx(request: Request):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    with SessionLocal() as db:
        rows = db.scalars(select(Student).order_by(Student.full_name)).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Participación"
    headers = ["DNI", "APELLIDOS Y NOMBRES", "GRADO", "SECCIÓN", "PARTICIPÓ", "FECHA Y HORA"]
    ws.append(headers)
    for r in rows:
        ws.append([r.dni or "", r.full_name, r.grade or "", r.section or "", "SÍ" if r.voted_at else "NO", r.voted_at.strftime("%Y-%m-%d %H:%M:%S") if r.voted_at else ""])

    fill = PatternFill("solid", fgColor="0F4279")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    widths = [14, 42, 10, 10, 12, 22]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=padron_participacion_elecciones.xlsx"},
    )


@app.get("/admin/export/results.csv")
def export_results(request: Request):
    if not admin_required(request):
        return RedirectResponse("/admin/login", status_code=303)
    if setting("election_open") == "1" and setting("show_results_live") != "1":
        return RedirectResponse("/admin?msg=Los+resultados+están+ocultos+mientras+la+elección+está+abierta", status_code=303)
    _, _, _, _, results, blank_result, _ = dashboard_data()
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["lista", "nombre_lista", "candidato", "votos", "porcentaje"])
    for r in results:
        w.writerow([r["list_number"], r["list_name"], r["candidate_name"], r["votes"], r["pct"]])
    w.writerow(["", "Voto en blanco", "", blank_result["votes"], blank_result["pct"]])
    data = out.getvalue().encode("utf-8-sig")
    return StreamingResponse(io.BytesIO(data), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=resultados_elecciones.csv"})


@app.get("/health")
def health():
    try:
        with SessionLocal() as db:
            db.execute(select(1))
        return {"status": "ok", "database": "postgresql" if "postgresql" in DATABASE_URL else "sqlite"}
    except Exception as e:
        return JSONResponse({"status": "error", "detail": str(e)}, status_code=500)
